"""Yuklangan fayllarni OpenAI Responses API tushunadigan bloklarga aylantirish.

PDF va rasmlar modelga o'z holicha yuboriladi — u ularni jadval va skanlar
bilan birga o'qiy oladi. Qolgan formatlardan matn lokal ajratib olinadi.
"""

from __future__ import annotations

import asyncio
import base64
import csv
import io
import json
from typing import Any

SUPPORTED_IMAGE_TYPES = {
    "image/jpeg": "image/jpeg",
    "image/png": "image/png",
    "image/gif": "image/gif",
    "image/webp": "image/webp",
}

TEXT_EXTENSIONS = {".txt", ".md", ".log", ".json", ".xml", ".html", ".csv", ".tsv"}

# Bitta hujjatdan olinadigan maksimal matn (taxminan 250k token'gacha sig'adi)
MAX_TEXT_CHARS = 400_000


class UnsupportedDocument(ValueError):
    """Fayl formati qo'llab-quvvatlanmaydi."""


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for i, table in enumerate(document.tables, start=1):
        parts.append(f"\n--- {i}-jadval ---")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"\n=== Varaq: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            parts.append(" | ".join("" if v is None else str(v) for v in row))
    wb.close()
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(presentation.slides, start=1):
        parts.append(f"\n--- {i}-slayd ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _extract_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    dialect = csv.Sniffer().sniff(text[:4096]) if text.strip() else csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return "\n".join(" | ".join(row) for row in reader)


def _extract_text(data: bytes, filename: str) -> str:
    if filename.lower().endswith((".csv", ".tsv")):
        try:
            return _extract_csv(data)
        except csv.Error:
            pass
    if filename.lower().endswith(".json"):
        try:
            return json.dumps(
                json.loads(data.decode("utf-8")), ensure_ascii=False, indent=2
            )
        except (ValueError, UnicodeDecodeError):
            pass
    return data.decode("utf-8", errors="replace")


def _build_sync(data: bytes, filename: str, mime: str) -> list[dict[str, Any]]:
    name = filename.lower()
    mime = (mime or "").lower()

    # 1. PDF — model o'zi o'qiydi (jadvallar, skanlar, grafiklar bilan)
    if mime == "application/pdf" or name.endswith(".pdf"):
        encoded = base64.standard_b64encode(data).decode("ascii")
        return [
            {
                "type": "input_file",
                "filename": filename if name.endswith(".pdf") else f"{filename}.pdf",
                "file_data": f"data:application/pdf;base64,{encoded}",
            }
        ]

    # 2. Rasm — chek, skrinshot, qo'lda yozilgan hisob
    media_type = SUPPORTED_IMAGE_TYPES.get(mime)
    if media_type is None and name.endswith((".jpg", ".jpeg")):
        media_type = "image/jpeg"
    elif media_type is None and name.endswith(".png"):
        media_type = "image/png"
    elif media_type is None and name.endswith(".webp"):
        media_type = "image/webp"
    if media_type:
        encoded = base64.standard_b64encode(data).decode("ascii")
        return [
            {
                "type": "input_image",
                "image_url": f"data:{media_type};base64,{encoded}",
                "detail": "high",  # chek va skanlardagi mayda raqamlar uchun
            }
        ]

    # 3. Office / matn formatlari — matnni lokal ajratamiz
    if name.endswith(".docx"):
        text = _extract_docx(data)
    elif name.endswith((".xlsx", ".xlsm")):
        text = _extract_xlsx(data)
    elif name.endswith(".pptx"):
        text = _extract_pptx(data)
    elif any(name.endswith(ext) for ext in TEXT_EXTENSIONS):
        text = _extract_text(data, filename)
    elif name.endswith((".doc", ".xls", ".ppt")):
        raise UnsupportedDocument(
            "Eski Office formati (.doc/.xls/.ppt) qo'llab-quvvatlanmaydi. "
            "Faylni .docx / .xlsx / .pptx yoki PDF ko'rinishida saqlab yuboring."
        )
    else:
        raise UnsupportedDocument(
            "Bu format qo'llab-quvvatlanmaydi.\n\n"
            "Qabul qilinadi: PDF, DOCX, XLSX, PPTX, CSV, TXT, JSON va rasmlar."
        )

    text = text.strip()
    if not text:
        raise UnsupportedDocument(
            "Fayldan matn topilmadi. Agar bu skan bo'lsa, PDF yoki rasm "
            "ko'rinishida yuboring — men uni o'qiy olaman."
        )

    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]

    note = "\n\n[DIQQAT: hujjat juda katta, faqat boshlang'ich qismi berildi]" if truncated else ""
    return [
        {
            "type": "input_text",
            "text": f"Hujjat nomi: {filename}\n\n<hujjat>\n{text}\n</hujjat>{note}",
        }
    ]


async def to_content_blocks(
    data: bytes, filename: str, mime: str = ""
) -> list[dict[str, Any]]:
    """Fayl baytlarini Responses API uchun content-bloklarga aylantiradi."""
    return await asyncio.to_thread(_build_sync, data, filename, mime)
