"""Tranzaksiyalarni formatlangan .xlsx faylga eksport qilish."""

from __future__ import annotations

import asyncio
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
INCOME_FILL = PatternFill("solid", fgColor="E2EFDA")
EXPENSE_FILL = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("Sana", "date", 12),
    ("Vaqt", "time", 10),
    ("Turi", "type", 10),
    ("Summa", "amount", 16),
    ("Valyuta", "currency", 10),
    ("Kategoriya", "category", 24),
    ("Tavsif", "description", 40),
    ("To'lov usuli", "payment_method", 14),
    ("Kontragent", "counterparty", 20),
    ("Manba", "source", 10),
]


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def _build_sync(
    transactions: list[dict[str, Any]], period_label: str, out_path: Path
) -> Path:
    wb = Workbook()

    # --- 1-varaq: tranzaksiyalar ------------------------------------------
    ws = wb.active
    ws.title = "Tranzaksiyalar"
    ws.append([title for title, _, _ in COLUMNS])
    _style_header(ws, len(COLUMNS))

    for idx, txn in enumerate(transactions, start=2):
        for col, (_, key, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=idx, column=col, value=txn.get(key, ""))
            cell.border = BORDER
            if key == "amount":
                cell.number_format = "#,##0.00"
                cell.font = Font(bold=True)
        fill = INCOME_FILL if txn.get("type") == "kirim" else EXPENSE_FILL
        ws.cell(row=idx, column=3).fill = fill

    for col, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    if transactions:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(transactions) + 1}"

    # --- 2-varaq: xulosa ---------------------------------------------------
    summary = wb.create_sheet("Xulosa")
    summary["A1"] = "MOLIYAVIY XULOSA"
    summary["A1"].font = Font(size=14, bold=True, color="1F4E78")
    summary["A2"] = f"Davr: {period_label}"
    summary["A3"] = f"Yaratildi: {datetime.now():%Y-%m-%d %H:%M}"
    summary["A4"] = f"Yozuvlar soni: {len(transactions)}"

    # Valyuta bo'yicha umumiy holat
    by_currency: dict[str, dict[str, float]] = defaultdict(
        lambda: {"kirim": 0.0, "chiqim": 0.0}
    )
    for txn in transactions:
        by_currency[txn.get("currency", config.DEFAULT_CURRENCY)][
            txn.get("type", "chiqim")
        ] += txn.get("amount", 0.0)

    row = 6
    summary.cell(row=row, column=1, value="Valyuta").font = Font(bold=True)
    summary.cell(row=row, column=2, value="Kirim").font = Font(bold=True)
    summary.cell(row=row, column=3, value="Chiqim").font = Font(bold=True)
    summary.cell(row=row, column=4, value="Balans").font = Font(bold=True)
    for col in range(1, 5):
        summary.cell(row=row, column=col).fill = HEADER_FILL
        summary.cell(row=row, column=col).font = HEADER_FONT

    for currency, totals in sorted(by_currency.items()):
        row += 1
        balance = totals["kirim"] - totals["chiqim"]
        summary.cell(row=row, column=1, value=currency)
        summary.cell(row=row, column=2, value=totals["kirim"]).number_format = "#,##0.00"
        summary.cell(row=row, column=3, value=totals["chiqim"]).number_format = "#,##0.00"
        cell = summary.cell(row=row, column=4, value=balance)
        cell.number_format = "#,##0.00"
        cell.font = Font(bold=True, color="1E7B34" if balance >= 0 else "C00000")

    # Kategoriyalar bo'yicha chiqim (asosiy valyuta)
    main_currency = config.DEFAULT_CURRENCY
    by_category: dict[str, float] = defaultdict(float)
    for txn in transactions:
        if txn.get("type") == "chiqim" and txn.get("currency") == main_currency:
            by_category[txn.get("category", "Boshqa chiqim")] += txn.get("amount", 0.0)

    cat_start = row + 3
    summary.cell(row=cat_start, column=1, value=f"Kategoriya bo'yicha chiqim ({main_currency})")
    summary.cell(row=cat_start, column=1).font = Font(bold=True, size=12)

    head = cat_start + 1
    summary.cell(row=head, column=1, value="Kategoriya").font = HEADER_FONT
    summary.cell(row=head, column=1).fill = HEADER_FILL
    summary.cell(row=head, column=2, value="Summa").font = HEADER_FONT
    summary.cell(row=head, column=2).fill = HEADER_FILL

    ordered = sorted(by_category.items(), key=lambda kv: kv[1], reverse=True)
    for i, (category, amount) in enumerate(ordered, start=1):
        summary.cell(row=head + i, column=1, value=category)
        summary.cell(row=head + i, column=2, value=amount).number_format = "#,##0.00"

    if ordered:
        last = head + len(ordered)
        pie = PieChart()
        pie.title = "Chiqimlar tarkibi"
        pie.height, pie.width = 9, 14
        pie.add_data(Reference(summary, min_col=2, min_row=head, max_row=last), titles_from_data=True)
        pie.set_categories(Reference(summary, min_col=1, min_row=head + 1, max_row=last))
        summary.add_chart(pie, "F6")

    # Oylar kesimi
    by_month: dict[str, dict[str, float]] = defaultdict(
        lambda: {"kirim": 0.0, "chiqim": 0.0}
    )
    for txn in transactions:
        if txn.get("currency") != main_currency:
            continue
        month = str(txn.get("date", ""))[:7]
        if month:
            by_month[month][txn.get("type", "chiqim")] += txn.get("amount", 0.0)

    if by_month:
        m_start = head + len(ordered) + 3
        summary.cell(row=m_start, column=1, value=f"Oylar kesimi ({main_currency})")
        summary.cell(row=m_start, column=1).font = Font(bold=True, size=12)
        m_head = m_start + 1
        for col, title in enumerate(["Oy", "Kirim", "Chiqim"], start=1):
            cell = summary.cell(row=m_head, column=col, value=title)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL

        for i, (month, totals) in enumerate(sorted(by_month.items()), start=1):
            summary.cell(row=m_head + i, column=1, value=month)
            summary.cell(row=m_head + i, column=2, value=totals["kirim"]).number_format = "#,##0.00"
            summary.cell(row=m_head + i, column=3, value=totals["chiqim"]).number_format = "#,##0.00"

        m_last = m_head + len(by_month)
        bar = BarChart()
        bar.type, bar.title = "col", "Kirim / Chiqim dinamikasi"
        bar.height, bar.width = 9, 18
        bar.add_data(
            Reference(summary, min_col=2, max_col=3, min_row=m_head, max_row=m_last),
            titles_from_data=True,
        )
        bar.set_categories(Reference(summary, min_col=1, min_row=m_head + 1, max_row=m_last))
        summary.add_chart(bar, "F26")

    for col, width in zip("ABCD", (28, 18, 18, 18)):
        summary.column_dimensions[col].width = width

    wb.save(out_path)
    return out_path


async def build_report(
    transactions: list[dict[str, Any]], period_label: str, user_id: int
) -> Path:
    """Excel hisobotini yaratadi va fayl yo'lini qaytaradi."""
    filename = f"hisobot_{user_id}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    out_path = config.EXPORT_DIR / filename
    return await asyncio.to_thread(_build_sync, transactions, period_label, out_path)
